from collections import Counter
from dataclasses import dataclass
from math import isfinite
from numbers import Real
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ActividadItem, Aplicacion, AplicacionActividad, ItemInstrumento


RUTA_OCUPACIONES = Path(__file__).resolve().parent / "archivos" / "Career_Interest_RIASEC_Clean.xlsx"


COLUMNAS_OCUPACIONES = ("code", "title", "R", "I", "A", "S", "E", "C")


@dataclass(frozen=True)
class OcupacionArchivo:
    codigo_onet: str
    titulo: str
    valores: tuple[float, ...]


def leer_ocupaciones(ruta: Path) -> list[OcupacionArchivo]:
    if not ruta.is_file():
        raise ValueError(f"Falta el archivo de ocupaciones: {ruta}")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException) as error:
        raise ValueError(f"No se puede leer el archivo de ocupaciones: {ruta}") from error
    try:
        filas = libro.worksheets[0].iter_rows(values_only=True)
        if tuple(next(filas, ())) != COLUMNAS_OCUPACIONES:
            raise ValueError("Columnas de ocupaciones inválidas; se esperan: code, title, R, I, A, S, E, C")
        ocupaciones = []
        codigos = set()
        for numero, fila in enumerate(filas, start=2):
            if all(valor is None for valor in fila):
                continue
            codigo, titulo, *valores = fila
            if not isinstance(codigo, str) or not codigo.strip():
                raise ValueError(f"Código O*NET inválido en la fila {numero}")
            codigo = codigo.strip()
            if codigo in codigos:
                raise ValueError(f"Código O*NET repetido: {codigo} (fila {numero})")
            if not isinstance(titulo, str) or not titulo.strip():
                raise ValueError(f"Título de ocupación inválido: {codigo} (fila {numero})")
            for dimension, valor in zip(COLUMNAS_OCUPACIONES[2:], valores, strict=True):
                if isinstance(valor, bool) or not isinstance(valor, Real) or not isfinite(valor):
                    raise ValueError(f"Valor no numérico o no finito en {codigo}, dimensión {dimension}, fila {numero}")
            ocupaciones.append(OcupacionArchivo(codigo, titulo.strip(), tuple(float(v) for v in valores)))
            codigos.add(codigo)
        if not ocupaciones:
            raise ValueError("El archivo de ocupaciones no contiene datos")
        return ocupaciones
    finally:
        libro.close()


def validar_distribucion_items(sesion: Session) -> None:
    """Cada aplicación presenta todos los ítems de su instrumento exactamente una vez."""
    for aplicacion in sesion.scalars(select(Aplicacion).order_by(Aplicacion.codigo)):
        esperados = dict(sesion.execute(select(ItemInstrumento.id, ItemInstrumento.codigo).where(
            ItemInstrumento.instrumento_id == aplicacion.instrumento_id,
        )).all())
        presentados = Counter(sesion.scalars(select(ActividadItem.item_id).join(
            AplicacionActividad, AplicacionActividad.actividad_id == ActividadItem.actividad_id,
        ).where(AplicacionActividad.aplicacion_id == aplicacion.id)))
        if not esperados or presentados != Counter({item: 1 for item in esperados}):
            faltantes = [codigo for item, codigo in esperados.items() if presentados[item] == 0]
            repetidos = [codigo for item, codigo in esperados.items() if presentados[item] > 1]
            ajenos = sesion.scalars(select(ItemInstrumento.codigo).where(
                ItemInstrumento.id.in_(set(presentados) - set(esperados)),
            )).all()
            raise ValueError(f"Distribución inválida en {aplicacion.codigo}: faltantes={faltantes}, "
                             f"repetidos={repetidos}, ajenos={ajenos}")
